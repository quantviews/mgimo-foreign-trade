import io
from unittest.mock import patch
from openpyxl import load_workbook
from app.comparison_data import export_comparison
from app.export_utils import build_export


def fixture():
    rows = [dict(strana=c, period=p, stoim=v, netto=v*2, type="fact", source="national")
            for c,p,v in [("IN","2026-06",100),("IN","2026-07",200),("CN","2026-06",300),("CN","2026-07",400),("CN","2026-08",900)]]
    rows.append(dict(strana="IN",period="2026-08",stoim=9999,netto=9999,type="pred",source="nowcast"))
    filters = dict(strana=["IN","CN"],napr=["ex"],tnved4=["2709"],metrics="stoim,netto",period_from="2026-06",period_to="2026-08")
    indices = [dict(strana=c,napr="ЭК",period=p,tn_level=4,tn_code="2709",idx=v,fizob_bp=1)
               for c,p,v in [("IN","2026-06",2),("IN","2026-07",3),("CN","2026-06",10),("CN","2026-07",11)]]
    return rows, filters, {"data":indices,"meta":{"has_more":False}}


def test_workbook_common_months_and_fizob():
    rows, filters, indices = fixture()
    version = {"period_max":"2026-08"}
    with patch("app.comparison_data.database_get", return_value=indices):
        comparison = export_comparison(rows,version,filters)
    assert comparison["common_fact_months"] == ["2026-06","2026-07"]
    assert comparison["excluded_months"] == ["2026-08"]
    assert comparison["totals"][0]["stoim"] == "300"
    assert comparison["physical_volume"]["series"][0]["change_percent"] == "50.0"
    content,_,_ = build_export(rows,version,filters,comparison)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Данные","Сравнение","Описание"]
    ws = wb["Сравнение"]
    assert ws["C2"].value == 300 / 10**9
    assert ws["J2"].value == 150
    assert ws["K2"].value == 50
    assert ws["L3"].value == -40
    assert ws["F2"].value is None
    assert ws["J2"].number_format.startswith("#,##0.00")
    assert wb["Данные"].max_row == len(rows)+1
    assert any("особенность и ценность" in str(r[1].value) for r in wb["Описание"])


def test_restricted_filters_never_expand_fizob():
    rows, filters, _ = fixture()
    filters["source"] = ["national"]
    with patch("app.comparison_data.database_get") as query:
        r = export_comparison(rows,{"period_max":"2026-08"},filters)
    query.assert_not_called()
    assert r["comparable"]
    assert not r["physical_volume"]["available"]


def test_missing_index_preserves_trade_totals():
    rows, filters, indices = fixture()
    indices["data"].pop()
    with patch("app.comparison_data.database_get",return_value=indices):
        r = export_comparison(rows,{"period_max":"2026-08"},filters)
    assert r["comparable"] and not r["physical_volume"]["available"]
    assert r["common_fact_months"] == ["2026-06","2026-07"]


def test_pred_only_and_ambiguous_direction():
    rows, filters, _ = fixture()
    r = export_comparison([r for r in rows if r["type"]=="pred"],{"period_max":"2026-08"},filters)
    assert not r["comparable"]
    filters["napr"] = ["im","ex"]
    assert not export_comparison(rows,{"period_max":"2026-08"},filters)["comparable"]
