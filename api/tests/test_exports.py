import asyncio
import csv
import io
from unittest.mock import AsyncMock, patch
from uuid import uuid4
import pytest
from fastapi import HTTPException
from app.routers import exports
from app.export_utils import build_export


def test_xlsx_all_rows_and_units():
    from openpyxl import load_workbook
    rows = [{'period': '2026-08-01', 'strana': 'CN', 'type': 'fact', 'tnved': '0101',
             'stoim': 2000000000, 'netto': 3000, 'label': '=BAD()'} for _ in range(80)]
    content, _, totals = build_export(rows, {'data_version': '2026-08'}, {'strana': ['CN']})
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ['Данные', 'Описание']
    ws = wb['Данные']
    assert ws.max_row == 81
    data = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert data['netto_tonnes'] == 3 and data['stoim_usd_bn'] == 2
    assert data['tnved'] == '0101' and data['label'] == '=BAD()'
    assert all(c.data_type != 'f' for row in ws for c in row)
    assert any('Фильтр: strana' == row[0].value for row in wb['Описание'])
    assert totals[('CN', 'fact')][0] == 160000000000


def test_latest_and_fixed():
    for mode, expected in [('latest', '2026-08-01'), ('fixed', '2025-12-01')]:
        with patch.object(exports.db, 'data_version', return_value={'period_max': '2026-08-01'}), \
             patch.object(exports.db, 'run_query', return_value=(['period','stoim','netto'], [('2025-01-01',10,20)])):
            _, result, _ = exports.prepare(exports.ExportQuery(strana=['CN'],period_from='2023-01',period_to='2025-12-01',end_mode=mode),100)
        assert result['filters']['period_to'] == expected
        assert 'type' in result['filters']['group_by']
        assert 'source' in result['filters']['group_by']
        assert result['meta']['complete']


def test_cap_does_not_publish_partial():
    with patch.object(exports.db,'data_version',return_value={'period_max':'2026-08-01'}), \
         patch.object(exports.db,'run_query',return_value=(['stoim'],[(1,),(2,)])):
        with pytest.raises(HTTPException) as error:
            exports.prepare(exports.ExportQuery(strana=['CN'],period_from='2023-01',),1)
    assert error.value.status_code == 413


def test_unrequested_weight_is_not_zero():
    with patch.object(exports.db,'data_version',return_value={'period_max':'2026-08-01'}), \
         patch.object(exports.db,'run_query',return_value=(['period','stoim'],[('2026-08-01',10)])):
        _, result, _ = exports.prepare(exports.ExportQuery(strana=['CN'],period_from='2023-01',metrics='stoim'),100)
    assert result['summary'][0]['netto'] is None


def test_download_checks_owner_and_expiry():
    async def run():
        pool = AsyncMock()
        pool.fetchrow.return_value = None
        with patch.object(exports.store,'init_pool',return_value=pool):
            with pytest.raises(HTTPException) as error:
                await exports.download_export(uuid4(), {'user_id':42})
        assert error.value.status_code == 404
        sql, _, uid = pool.fetchrow.call_args.args
        assert 'user_id=$2' in sql and 'expires_at>now()' in sql and uid == 42
    asyncio.run(run())


def test_export_filters_required_and_period_bounded():
    cases = [{}, {'strana':['CN']}, {'strana':['CN'], 'period_from':'2019-01'},
             {'strana':['CN'], 'period_from':'2027-01'}]
    for args in cases:
        with patch.object(exports.db, 'data_version', return_value={'period_max':'2026-08-01'}), patch.object(exports.db, 'run_query') as query:
            with pytest.raises(HTTPException) as exc:
                exports.prepare(exports.ExportQuery(**args), 1000000)
            assert exc.value.status_code == 400
            query.assert_not_called()


def test_global_row_cap_even_for_large_plan():
    with patch.object(exports.db, 'data_version', return_value={'period_max':'2026-08-01'}), patch.object(exports.db, 'run_query', return_value=(['stoim'], [(1,)] * 1001)):
        with pytest.raises(HTTPException) as exc:
            exports.prepare(exports.ExportQuery(strana=['CN'],period_from='2026-01'), 1000000)
        assert exc.value.status_code == 413


def test_export_delegation_is_not_client_controlled():
    from types import SimpleNamespace
    async def run():
        pool = AsyncMock()
        user = {"user_id": 42, "token_id": 5}
        assert await exports.export_owner(SimpleNamespace(headers={}), user, pool) == 42
        pool.fetchval.assert_not_called()
        pool.fetchval.return_value = False
        with pytest.raises(HTTPException) as exc:
            await exports.export_owner(SimpleNamespace(headers={"x-export-user-id": "7"}), user, pool)
        assert exc.value.status_code == 403
        pool.fetchval.side_effect = [True, True]
        assert await exports.export_owner(SimpleNamespace(headers={"x-export-user-id": "7"}), user, pool) == 7
    asyncio.run(run())


def test_quota_rejection_never_saves_or_returns_inline_data():
    from unittest.mock import MagicMock
    from types import SimpleNamespace
    async def run():
        conn = AsyncMock()
        conn.fetchval.return_value = 3600
        conn.transaction = MagicMock(return_value=AsyncMock())
        pool = MagicMock()
        pool.acquire.return_value.__aenter__.return_value = conn
        request = SimpleNamespace(headers={}, state=SimpleNamespace())
        with patch.object(exports.store, 'init_pool', return_value=pool), patch.object(exports, 'prepare', return_value=(b'xlsx', {}, [{}])):
            with pytest.raises(HTTPException) as exc:
                await exports.create_export(exports.ExportQuery(inline=True), request, {'user_id':42, 'plan':{'max_rows':1000}})
        assert exc.value.status_code == 429
        assert exc.value.headers['Retry-After'] == '3600'
        assert conn.fetchval.call_count == 1
        conn.execute.assert_not_called()
    asyncio.run(run())

def test_yoy_windows_separate_sources_and_global_cap():
    from openpyxl import load_workbook
    columns = ['period','strana','type','source','stoim','netto','kol','tnved6']
    previous = [('2024-09-01','CN','fact','national',100,None,2,'870360')]
    current = [('2025-09-01','CN','fact','national',150,1000,3,'870360'),
               ('2025-09-01','CN','pred','nowcast',0.115,None,None,'870360')]
    body = exports.ExportQuery(strana=['CN'], period_from='2025-09',period_to='2025-11',
                               end_mode='fixed',compare_yoy=True)
    with patch.object(exports.db,'data_version',return_value={'period_max':'2026-08-01'}), \
         patch.object(exports.db,'run_query',side_effect=[(columns,previous),(columns,current)]), \
         patch.object(exports,'build_trade_query', wraps=exports.build_trade_query) as query:
        content,result,rows = exports.prepare(body,1000)
    assert [(c.kwargs['period_from'],c.kwargs['period_to']) for c in query.call_args_list] == [
        ('2024-09-01','2024-11-01'),('2025-09-01','2025-11-01')]
    assert result['meta']['rows'] == 3
    assert result['file_manifest']['sheets'] == {'Факт':2,'Nowcast':1}
    wb=load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ['Факт','Nowcast','Описание']
    def records(sheet):
        data=list(wb[sheet].values)
        return [dict(zip(data[0],r)) for r in data[1:]]
    facts=records('Факт'); estimates=records('Nowcast')
    assert {r['comparison_period'] for r in facts} == {'current','previous'}
    assert facts[0]['netto'] is None
    assert estimates[0]['stoim'] == 0.115
    assert all(r['type']=='fact' for r in facts)
    with patch.object(exports.db,'data_version',return_value={'period_max':'2026-08-01'}), \
         patch.object(exports.db,'run_query',side_effect=[(columns,previous),(columns,current)]):
        with pytest.raises(HTTPException) as error: exports.prepare(body,2)
    assert error.value.status_code==413

def test_source_estimates_are_not_project_nowcast():
    from openpyxl import load_workbook
    rows=[{'period':'2026-07-01','type':'fact','source':'comtrade','kol':0.115,
           'kol_estimated':True,'quality_flags':'SOURCE_QTY_ESTIMATED','comparison_period':'current'},
          {'period':'2026-07-01','type':'pred','source':'nowcast','kol':None,'comparison_period':'current'}]
    content,_,_=build_export(rows,{}, {'compare_yoy':True,'yoy_periods':{'current':['2026-07','2026-07'],'previous':['2025-07','2025-07']}})
    wb=load_workbook(io.BytesIO(content))
    values=list(wb['Факт'].values);r=dict(zip(values[0],values[1]))
    assert r['source']=='comtrade' and r['kol']==0.115 and r['kol_estimated'] is True
    assert r['quality_flags']=='SOURCE_QTY_ESTIMATED'
    assert wb['Факт'].max_row==2 and wb['Nowcast'].max_row==2
