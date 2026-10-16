"""Bounded Excel export construction for the production API."""
from decimal import Decimal
from datetime import datetime, timezone
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def text_cell(cell, value):
    cell.value = str(value)
    cell.data_type = "s"  # Codes retain leading zeros; untrusted text is never a formula.


def build_export(rows, meta, filters, comparison=None):
    normalized = [{k.lower(): v for k, v in row.items()} for row in rows]
    columns = list(dict.fromkeys(k for row in normalized for k in row))
    headers = columns + ["stoim_usd_bn", "netto_tonnes", "data_version"]
    wb = Workbook()
    ws = wb.active
    separate = bool(filters.get("compare_yoy"))
    ws.title = "Факт" if separate else "Данные"
    ws.append(headers)
    fact_sheet = ws
    estimate_sheet = wb.create_sheet("Nowcast") if separate else None
    if estimate_sheet is not None: estimate_sheet.append(headers)
    totals = {}
    for row in normalized:
        if separate:
            ws = estimate_sheet if row.get("type") == "pred" or row.get("source") == "nowcast" else fact_sheet
        key = (row.get("strana", "ALL"), row.get("type", "unspecified"))
        sums = totals.setdefault(key, [Decimal(0), Decimal(0)])
        value = Decimal(str(row["stoim"])) if row.get("stoim") is not None else None
        weight = Decimal(str(row["netto"])) if row.get("netto") is not None else None
        if value is not None: sums[0] += value
        if weight is not None: sums[1] += weight
        values = [row.get(c) for c in columns] + [value / 10**9 if value is not None else None, weight / 1000 if weight is not None else None, meta.get("data_version", "")]
        idx = ws.max_row + 1
        for col, val in enumerate(values, 1):
            cell = ws.cell(idx, col)
            if isinstance(val, str): text_cell(cell, val)
            else: cell.value = val
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if comparison is not None:
        add_comparison(wb, comparison)
    info = wb.create_sheet("Описание")
    details = [
        ("Параметр", "Описание"),
        ("Создано (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Версия данных", meta.get("data_version", "")),
        ("Последний период базы", meta.get("period_max", "")),
        ("Строк данных", len(rows)),
        ("Источники в выгрузке", ", ".join(sorted({str(r.get("source", "не указан")) for r in normalized}))),
        ("national", "Национальная статистика стран-партнеров; методология: https://nts.mgimo.ru/site/technical.html"),
        ("comtrade", "UN Comtrade: https://comtradeplus.un.org/"),
        ("nowcast", "Модельная оценка проекта, не фактическая таможенная статистика"),
        ("TYPE", "fact - факт; pred - модельная оценка"),
        ("STOIM / stoim_usd_bn", "Доллары США / млрд долларов США, текущие цены"),
        ("NETTO / netto_tonnes", "Килограммы / тонны; масса не равна индексу физического объема"),
        ("KOL", "Количество в EDIZM / EDIZM_ISO. Разные единицы нельзя складывать"),
        ("Оценки источника", "NETTO_ESTIMATED / KOL_ESTIMATED: true - показатель оценен источником, false - источник не пометил его как оценку, пусто - признак неизвестен. SOURCE_QTY_ESTIMATED / SOURCE_NETTO_ESTIMATED сохраняют исходные флаги Comtrade. TYPE=fact не исключает оценки физических показателей источником; это отдельно от nowcast проекта."),
        ("SOURCE_UNIT", "Единица количества до общей нормализации. Неоднозначные обозначения сохраняются без предположений о пересчете."),
        ("quality_flags", "Предупреждения источника: неизвестная масса, скрытая детализация, подозрительно малая масса транспортного средства. Это сигнал для проверки, не автоматическое исправление значений."),
        ("Неизвестная масса", "Пустая ячейка NETTO не равна нулю. Если хотя бы одна составляющая массы неизвестна, агрегированная масса не публикуется как полный итог."),
        ("Коды товаров", "Текстовые значения с сохранением ведущих нулей"),
        ("Полнота", "Все строки выбранного среза в пределах лимита. Отсутствующие наблюдения не равны нулю; покрытие источников может различаться"),
        ("Ограничения", "До 1000 строк (или меньше по тарифу), 60 месяцев, 5 МБ. Выбор стран или товарной группы обязателен. Ссылка действует 7 дней"),
    ]
    if separate:
        details += [
            ("Два периода", json.dumps(filters["yoy_periods"], ensure_ascii=False)),
            ("comparison_period", "current - текущий период; previous - те же месяцы годом ранее. Оба периода сохранены для воспроизведения расчета год к году."),
            ("Листы", "Факт - статистика источников с признаками оценок Comtrade. Nowcast - только модельные оценки проекта, не входят в фактические итоги. Пустой лист означает отсутствие строк в выбранном срезе, не нулевые поставки."),
            ("Расчеты", "Исходные помесячные строки по подкатегориям. Доли и вклады в динамику рассчитываются в аналитическом ответе; в файле они не добавлены. Не складывайте текущий и прошлогодний периоды в один итог."),
        ]
    if comparison is not None:
        details += [
            ('Сравнение: статус', 'Рассчитано' if comparison.get('comparable') else comparison['answer_ru']),
            ('Общие месяцы факта', ', '.join(comparison.get('common_fact_months', []))),
            ('Исключенные месяцы', ', '.join(comparison.get('excluded_months', []))),
            ('Выбор периода', 'Пересечение месяцев с ненулевыми по наличию (не NULL) стоимостью и массой у всех выбранных стран внутри запрошенного окна. Числовой ноль допустим. Окно назад не сдвигается. pred исключен только из сравнения, исходный лист сохраняет выбранные данные.'),
            ('Разницы и проценты', 'База - первая выбранная страна. Разница = страна минус база; процент = разница / база × 100. При неположительной базе процент не рассчитан. Пустая ячейка не означает ноль.'),
            ('Стоимость и масса', 'Суммы только за одинаковые перечисленные месяцы. Стоимость в текущих USD, масса в кг; в сравнении млрд USD и млн тонн.'),
            ('fizob: методология', 'Индекс проекта использует 12-месячное сглаживание и восстановление части количеств. В товарных агрегатах применяются меняющиеся стоимостные веса. Агрегация товаров с разными физическими единицами в страновом итоге - особенность и ценность методологии.'),
            ('fizob: расчет динамики', '(idx конца / idx начала - 1) × 100. Каждая страна нормирована на собственное начало = 100. Это динамика сглаженного индекса между крайними общими месяцами, не рост суммы физических объемов за весь период и не обязательно год к году. Индексы не складываются; абсолютные уровни разных стран не сравниваются.'),
            ('fizob: сопоставимость', 'Проверяются наличие индекса за каждый общий месяц и постоянство знаменателя нормировки. Разница темпов между странами указана в процентных пунктах. Масса не равна fizob; точное разложение стоимости на цену и объем не выполняется.'),
            ('fizob: статус', 'Рассчитано' if comparison.get('physical_volume', {}).get('available') else comparison.get('physical_volume', {}).get('reason', 'Нет сопоставимого периода.')),
            ('Сопоставимость покрытия', 'Общие месяцы не гарантируют полноты товарного покрытия. Сравнение рассчитано по строкам листа «Данные» с теми же фильтрами; фильтры ради индекса не расширяются.'),
            ('Методология проекта', 'https://nts.mgimo.ru/site/technical.html'),
        ]
    details += [("Фильтр: " + k, json.dumps(v, ensure_ascii=False)) for k, v in filters.items() if k != "inline"]
    for a, b in details:
        info.append([a, None])
        text_cell(info.cell(info.max_row, 2), b)
    from openpyxl.styles import Alignment
    for cell in info["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        info.row_dimensions[cell.row].height = max(30, 15 * (len(str(cell.value or '')) // 85 + 1))
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 105
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16324F")
    for sheet in ([fact_sheet, estimate_sheet] if separate else [ws]):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), normalized, totals


def add_comparison(wb, result):
    from openpyxl.styles import Alignment
    sheet = wb.create_sheet('Сравнение')
    headers = ['Страна', 'Код', 'Стоимость, млрд USD', 'Масса, млн т',
               'Разница стоимости, млрд USD', 'Разница стоимости, %',
               'Разница массы, млн т', 'Разница массы, %',
               'fizob: начало = 100', 'fizob: конец', 'Динамика fizob, %', 'Разница динамики fizob, п.п.']
    sheet.append(headers)
    if result.get('comparable'):
        differences = {r['country']:r for r in result['differences']}
        physical = result.get('physical_volume', {})
        indices = {r['strana']:r for r in physical.get('series', [])}
        index_diffs = {r['country']:r['growth_difference_pp'] for r in physical.get('differences', [])}
        def decimal(v, scale=1):
            return Decimal(str(v)) / scale if v is not None else None
        for row in result['totals']:
            country = row['strana']; diff = differences.get(country, {})
            idx = indices.get(country, {})
            sheet.append([row['country_name'], country, decimal(row['stoim'], 10**9), decimal(row['netto'], 10**9),
                          decimal(diff.get('stoim', {}).get('difference'), 10**9), decimal(diff.get('stoim', {}).get('percent')),
                          decimal(diff.get('netto', {}).get('difference'), 10**9), decimal(diff.get('netto', {}).get('percent')),
                          decimal(idx.get('rebased_start')), decimal(idx.get('rebased_end')), decimal(idx.get('change_percent')), decimal(index_diffs.get(country))])
        sheet.auto_filter.ref = sheet.dimensions
        notes = [('База сравнения', result['baseline_country']),
                 ('Общие месяцы факта', ', '.join(result['common_fact_months'])),
                 ('Исключенные месяцы', ', '.join(result['excluded_months']) or 'Нет'),
                 ('fizob', 'Динамика между первым и последним общим месяцем; методология на листе «Описание».' if physical.get('available') else physical.get('reason', 'Не рассчитан'))]
    else:
        notes = [('Сравнение не рассчитано', result['answer_ru'])]
    sheet.append([])
    for key, value in notes:
        sheet.append([key, value])
        sheet.merge_cells(start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=12)
        sheet.row_dimensions[sheet.max_row].height = 44
    sheet.freeze_panes = 'C2'
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str): text_cell(cell, cell.value)
            elif cell.value is not None: cell.number_format = '#,##0.00;[Red]-#,##0.00'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[1].height = 60
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 23
